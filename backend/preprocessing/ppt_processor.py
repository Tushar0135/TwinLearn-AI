"""
PPT/PPTX lecture processor.

Extracts meaningful text slide-by-slide while preserving
slide structure.
"""

from pathlib import Path

from pptx import Presentation


class PPTProcessingError(Exception):
    """Raised when a PowerPoint file cannot be processed."""
    pass


def extract_text_from_shape(shape) -> list[str]:
    """
    Extract text from a PowerPoint shape.

    Handles normal text boxes, titles, subtitles, and
    other shapes containing text.
    """

    extracted = []

    if not hasattr(shape, "text_frame"):
        return extracted

    if shape.text_frame is None:
        return extracted

    for paragraph in shape.text_frame.paragraphs:
        text = paragraph.text.strip()

        if text:
            extracted.append(text)

    return extracted


def process_pptx(file_path: str | Path) -> dict:
    """
    Extract text from a PPTX file while preserving slide boundaries.

    Returns:
        {
            "pages_or_slides": int,
            "text": str
        }

    Raises:
        PPTProcessingError:
            If the presentation cannot be opened or processed.
    """

    path = Path(file_path)

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise PPTProcessingError(
            f"Unable to open PowerPoint file: {exc}"
        ) from exc

    slides = []

    try:
        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):
            slide_lines = []

            for shape in slide.shapes:
                shape_text = extract_text_from_shape(shape)
                slide_lines.extend(shape_text)

            slide_text = "\n".join(slide_lines).strip()

            slides.append(
                f"[Slide {slide_number}]\n{slide_text}"
            )

    except Exception as exc:
        raise PPTProcessingError(
            f"Failed while extracting PowerPoint text: {exc}"
        ) from exc

    if not slides:
        raise PPTProcessingError(
            "PowerPoint presentation contains no slides."
        )

    return {
        "pages_or_slides": len(slides),
        "text": "\n\n".join(slides)
    }


def process_ppt(file_path: str | Path) -> dict:
    """
    Process a PowerPoint file.

    Currently supports PPTX directly.

    Legacy .ppt support will be added separately because
    python-pptx does not support the old binary .ppt format.
    """

    path = Path(file_path)

    if path.suffix.lower() == ".ppt":
        raise PPTProcessingError(
            "Legacy .ppt files require a separate conversion/extraction "
            "method. Direct .ppt extraction is not supported by "
            "python-pptx."
        )

    return process_pptx(path)

def process_ppt(file_path: str) -> dict:
    """
    Handle legacy .ppt files.

    python-pptx supports .pptx but not the old binary .ppt format.
    Therefore, legacy PPT files receive a clear standardized error
    instead of being incorrectly processed as PPTX.
    """

    path = Path(file_path)

    raise PPTProcessingError(
        f"Legacy .ppt format is not directly supported by python-pptx: "
        f"{path.name}. Please convert the file to .pptx before uploading."
    )